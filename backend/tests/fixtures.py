"""Test fixtures that build real files, so extraction is tested end to end."""

from __future__ import annotations

import io


def make_pdf(pages: list[str]) -> bytes:
    """A minimal, valid PDF with one text page per entry (Helvetica, one line per row)."""
    objects: list[bytes] = []

    def add(body: bytes) -> int:
        objects.append(body)
        return len(objects)

    font = add(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    page_ids: list[int] = []
    pages_id = len(objects) + 1 + 2 * len(pages)  # reserved below
    for text in pages:
        lines = [ln.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)") for ln in text.splitlines()]
        ops = ["BT", "/F1 11 Tf", "14 TL", "50 780 Td"]
        for line in lines:
            ops.append(f"({line}) Tj T*")
        ops.append("ET")
        stream = "\n".join(ops).encode("latin-1", errors="replace")
        content = add(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
        page = add(
            b"<< /Type /Page /Parent %d 0 R /MediaBox [0 0 595 842] /Contents %d 0 R "
            b"/Resources << /Font << /F1 %d 0 R >> >> >>" % (pages_id, content, font)
        )
        page_ids.append(page)
    kids = b" ".join(b"%d 0 R" % p for p in page_ids)
    assert add(b"<< /Type /Pages /Kids [%s] /Count %d >>" % (kids, len(page_ids))) == pages_id
    catalog = add(b"<< /Type /Catalog /Pages %d 0 R >>" % pages_id)

    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % i + body + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1))
    for off in offsets:
        out.write(b"%010d 00000 n \n" % off)
    out.write(b"trailer\n<< /Size %d /Root %d 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, catalog, xref))
    return out.getvalue()


def make_docx(paragraphs: list[tuple[str, str | None]]) -> bytes:
    """``paragraphs`` are ``(text, style)``; style ``"Heading 1"`` etc. or ``None``."""
    import docx

    document = docx.Document()
    for text, style in paragraphs:
        if style:
            document.add_paragraph(text, style=style)
        else:
            document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


INSTRUCTION_TEXT = """PANDUAN PENULISAN SKRIPSI
Program Studi Informatika
1. Ketentuan Umum
Skripsi wajib menggunakan minimal 20 referensi dari 5 tahun terakhir.
Format sitasi harus mengikuti gaya IEEE.
Naskah paling sedikit 40 halaman tidak termasuk lampiran.
2. Jadwal Penting
Batas pengumpulan proposal: 30 Oktober 2026
Seminar proposal dilaksanakan pada 20 November 2026
Sidang akhir paling lambat 15 Maret 2027
"""

JOURNAL_PAGES = [
    "Decision Making for NPCs in Real-Time Strategy Games\nAbstract\n"
    "This paper compares Behavior Trees, Utility AI and Monte Carlo Tree Search (MCTS)\n"
    "for non-player character decisions. doi: 10.1234/rts.2024.001\nKeywords: RTS, MCTS, game AI",
    "1. Introduction\nReal-time strategy games need fast decisions under uncertainty.\n"
    "2. Method\nWe implement MCTS with a 50 ms budget and compare it with Utility AI.\n"
    "MCTS explores possible futures by random playouts and keeps the most promising moves.",
    "3. Results\nMCTS won 68 percent of matches against the Utility AI baseline.\n"
    "Behavior Trees were fastest but least adaptive.\n4. Conclusion\n"
    "MCTS offers the best balance for tactical decisions in RTS games.",
]
