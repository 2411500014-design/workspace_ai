"""Text extraction per page (master plan §9).

PDF uses pypdf (BSD) and DOCX uses python-docx (MIT); PyMuPDF is avoided because of its
AGPL licence. DOCX and plain text have no real pages, so they are split into sections of
roughly one page each and numbered the same way.
"""

from __future__ import annotations

import io
from dataclasses import dataclass

from app.core.errors import AppError

PSEUDO_PAGE_CHARS = 3000

SUPPORTED = {
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "text/plain": "txt",
    "text/markdown": "md",
}
EXTENSIONS = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".txt": "text/plain",
    ".md": "text/markdown",
}


@dataclass(frozen=True)
class Page:
    number: int  # 1-based
    text: str


def detect_mime(filename: str, declared: str | None) -> str:
    lower = filename.lower()
    for ext, mime in EXTENSIONS.items():
        if lower.endswith(ext):
            return mime
    if declared in SUPPORTED:
        return declared
    raise AppError("unsupported_file_type", 415, {"filename": filename})


def extract_pages(data: bytes, mime: str) -> list[Page]:
    kind = SUPPORTED.get(mime)
    if kind == "pdf":
        return _pdf(data)
    if kind == "docx":
        return _docx(data)
    if kind in ("txt", "md"):
        return _paginate(_decode(data))
    raise AppError("unsupported_file_type", 415, {"mime": mime})


def count_pages(data: bytes, mime: str) -> int:
    if SUPPORTED.get(mime) == "pdf":
        from pypdf import PdfReader

        try:
            return len(PdfReader(io.BytesIO(data)).pages)
        except Exception as exc:  # corrupt or encrypted PDF
            raise AppError("unreadable_file", 422) from exc
    return len(extract_pages(data, mime))


def _pdf(data: bytes) -> list[Page]:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception as exc:
                raise AppError("encrypted_file", 422) from exc
        return [Page(i + 1, (page.extract_text() or "").strip()) for i, page in enumerate(reader.pages)]
    except AppError:
        raise
    except Exception as exc:
        raise AppError("unreadable_file", 422) from exc


def _docx(data: bytes) -> list[Page]:
    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise AppError("unreadable_file", 422) from exc
    lines = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = (paragraph.style.name or "") if paragraph.style is not None else ""
        # Mark headings so the chunker can follow the document structure.
        lines.append(f"# {text}" if style.lower().startswith(("heading", "judul", "title")) else text)
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                lines.append(" | ".join(cells))
    return _paginate("\n".join(lines))


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _paginate(text: str) -> list[Page]:
    pages: list[Page] = []
    buffer: list[str] = []
    size = 0
    for line in text.splitlines():
        buffer.append(line)
        size += len(line) + 1
        if size >= PSEUDO_PAGE_CHARS:
            pages.append(Page(len(pages) + 1, "\n".join(buffer).strip()))
            buffer, size = [], 0
    if buffer and "\n".join(buffer).strip():
        pages.append(Page(len(pages) + 1, "\n".join(buffer).strip()))
    return pages or [Page(1, "")]
